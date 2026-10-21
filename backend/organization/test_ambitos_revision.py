"""Regresiones de revisión interna por parejas unidad-período y revocación."""

from datetime import timedelta
from io import BytesIO

from django.contrib.auth.models import Group, Permission
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Usuario
from auditoria.models import Auditoria
from catalogo.models import Dimension, VersionMarcoEvaluacion
from evaluation.models import Asignacion, Criterio, HistorialEstado, Indicador, Periodo
from evidence.models import Evidencia, Observacion, VersionEvidencia
from organization.models import (
    AmbitoEvaluacion,
    AmbitoRevision,
    PerfilUsuario,
    TipoUnidadOrganizacional,
    UnidadOrganizacional,
)
from vinculacion.models import Correspondencia


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AmbitosRevisionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        hoy = timezone.localdate()
        cls.p1 = Periodo.objects.create(nombre="P1 revisión", fecha_inicio=hoy, fecha_fin=hoy + timedelta(days=365))
        cls.p2 = Periodo.objects.create(nombre="P2 revisión", fecha_inicio=hoy, fecha_fin=hoy + timedelta(days=730), activo=False)
        cls.p3 = Periodo.objects.create(nombre="Sin autorización", fecha_inicio=hoy, fecha_fin=hoy + timedelta(days=730), activo=False)
        facultad, _ = TipoUnidadOrganizacional.objects.get_or_create(nombre="Facultad")
        centro, _ = TipoUnidadOrganizacional.objects.get_or_create(nombre="Centro")
        unidad = TipoUnidadOrganizacional.objects.create(nombre="Unidad de prueba")
        cls.raiz_a = UnidadOrganizacional.objects.create(nombre="Facultad A", tipo=facultad)
        cls.raiz_b = UnidadOrganizacional.objects.create(nombre="Centro B", tipo=centro)
        cls.a = UnidadOrganizacional.objects.create(nombre="Unidad A", tipo=unidad, unidad_padre=cls.raiz_a)
        cls.b = UnidadOrganizacional.objects.create(nombre="Unidad B", tipo=unidad, unidad_padre=cls.raiz_b)
        cls.nieta = UnidadOrganizacional.objects.create(nombre="Subunidad A", tipo=unidad, unidad_padre=cls.a)

        def usuario(nombre, rol):
            user = Usuario.objects.create_user(username=nombre, email=f"{nombre}@example.test", password="testpass")
            group, _ = Group.objects.get_or_create(name=rol)
            user.groups.add(group)
            return user

        cls.admin = usuario("admin_revision", "Administrador General")
        cls.coordinador = usuario("coord_revision", "Coordinador Quinquenal")
        cls.revisor = usuario("revisor_explicito", "Revisor Institucional")
        cls.sin_ambitos = usuario("revisor_sin_ambitos", "Revisor Institucional")
        cls.externo = usuario("externo_revision", "Evaluador Externo")
        cls.responsable = usuario("responsable_revision", "Responsable Departamental")
        PerfilUsuario.objects.create(usuario=cls.revisor, unidad_organizacional=cls.b)
        PerfilUsuario.objects.create(usuario=cls.sin_ambitos, unidad_organizacional=cls.a)
        cls.ambito_a = AmbitoRevision.objects.create(usuario=cls.revisor, unidad_organizacional=cls.raiz_a, periodo=cls.p1)
        cls.ambito_b = AmbitoRevision.objects.create(usuario=cls.revisor, unidad_organizacional=cls.raiz_b, periodo=cls.p2)
        AmbitoEvaluacion.objects.create(usuario=cls.externo, unidad_organizacional=cls.raiz_a, periodo=cls.p1)
        criterio = Criterio.objects.create(nombre="Referencia compartida", periodo=cls.p1)
        cls.indicador = Indicador.objects.create(nombre="Indicador revisión", criterio=criterio)
        cls.asignaciones = {}
        cls.evidencias = {}
        cls.versiones = {}
        for letra, unidad_obj in (("a", cls.a), ("b", cls.b)):
            for numero, periodo in ((1, cls.p1), (2, cls.p2)):
                key = f"{letra}{numero}"
                asignacion = Asignacion.objects.create(indicador=cls.indicador, unidad_responsable=unidad_obj, periodo=periodo, estado="en_progreso")
                evidencia = Evidencia.objects.create(titulo=key, descripcion="Prueba", asignacion=asignacion)
                version = VersionEvidencia.objects.create(
                    evidencia=evidencia,
                    archivo=SimpleUploadedFile(f"{key}.txt", key.encode(), content_type="text/plain"),
                    subido_por=cls.responsable,
                )
                cls.asignaciones[key] = asignacion
                cls.evidencias[key] = evidencia
                cls.versiones[key] = version
        cls.obs_ajena = Observacion.objects.create(version=cls.versiones["b1"], usuario=cls.admin, comentario="Ajena")

    def setUp(self):
        cache.clear()
        self.client = APIClient()
        self.client.force_authenticate(self.revisor)

    def ids(self, response, campo="id"):
        self.assertEqual(response.status_code, 200, response.data)
        return {item[campo] for item in response.data["results"]}

    def test_listados_no_cruzan_unidades_y_periodos(self):
        for p, key in ((self.p1, "a1"), (self.p2, "b2")):
            with self.subTest(periodo=p.nombre):
                self.assertEqual(self.ids(self.client.get(f"/api/asignaciones/?periodo={p.pk}")), {self.asignaciones[key].pk})
                self.assertEqual(self.ids(self.client.get(f"/api/evidencias/?periodo={p.pk}"), "id_evidencia"), {self.evidencias[key].pk})
                self.assertEqual(self.ids(self.client.get(f"/api/versiones/?periodo={p.pk}"), "id_version"), {self.versiones[key].pk})

    def test_alcance_incluye_raiz_y_descendientes_sin_depender_del_perfil(self):
        for unidad in (self.raiz_a, self.nieta):
            asignacion = Asignacion.objects.create(indicador=self.indicador, unidad_responsable=unidad, periodo=self.p1)
            self.assertEqual(self.client.get(f"/api/asignaciones/{asignacion.pk}/").status_code, 200)

    def test_detalles_y_archivos_ajenos_no_son_accesibles(self):
        for key in ("a2", "b1"):
            p = self.asignaciones[key].periodo_id
            urls = [
                f"/api/asignaciones/{self.asignaciones[key].pk}/?periodo={p}",
                f"/api/evidencias/{self.evidencias[key].pk}/detalle/?periodo={p}",
                f"/api/evidencias/{self.evidencias[key].pk}/historial/?periodo={p}",
                f"/api/versiones/{self.versiones[key].pk}/descargar/?periodo={p}",
                f"/api/versiones/{self.versiones[key].pk}/preview/?periodo={p}",
            ]
            for url in urls:
                with self.subTest(url=url):
                    self.assertEqual(self.client.get(url).status_code, 404)

    def test_descarga_dentro_del_centro(self):
        response = self.client.get(f"/api/versiones/{self.versiones['b2'].pk}/descargar/?periodo={self.p2.pk}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(response.streaming_content), b"b2")

    def test_zip_rechaza_seleccion_mixta(self):
        response = self.client.post("/api/evidencias/descargas-masivas/", {"evidencias": [self.evidencias["a1"].pk, self.evidencias["b1"].pk]}, format="json")
        self.assertEqual(response.status_code, 404)

    def test_revisor_puede_observar_y_aprobar_en_su_ambito(self):
        asig = self.asignaciones["a1"]
        response = self.client.post(f"/api/asignaciones/{asig.pk}/observada/", {"comentario": "Corregir"}, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertTrue(Observacion.objects.filter(version=self.versiones["a1"], usuario=self.revisor, comentario="Corregir").exists())
        self.assertTrue(HistorialEstado.objects.filter(asignacion=asig, estado_nuevo="observada", usuario=self.revisor).exists())
        # Otra entrega disponible para revisión en la unidad raíz autorizada.
        asig2 = Asignacion.objects.create(indicador=self.indicador, unidad_responsable=self.raiz_a, periodo=self.p1, estado="en_progreso")
        response = self.client.post(f"/api/asignaciones/{asig2.pk}/aprobar/", {}, format="json")
        self.assertEqual(response.status_code, 200)

    def test_revision_ajena_no_cambia_estado_ni_historial(self):
        asig = self.asignaciones["b1"]
        for accion in ("aprobar", "rechazar", "observada", "en_revision"):
            self.assertEqual(self.client.post(f"/api/asignaciones/{asig.pk}/{accion}/", {"comentario": "Ajeno"}, format="json").status_code, 404)
        asig.refresh_from_db()
        self.assertEqual(asig.estado, "en_progreso")
        self.assertFalse(HistorialEstado.objects.filter(asignacion=asig).exists())

    def test_observacion_directa_valida_version_recibida(self):
        for key, esperado in (("b1", 400), ("a1", 201)):
            response = self.client.post("/api/observaciones/", {"version": self.versiones[key].pk, "comentario": "Revisar"}, format="json")
            self.assertEqual(response.status_code, esperado, response.data)

    def test_no_se_puede_reasociar_observacion_a_version_ajena(self):
        obs = Observacion.objects.create(version=self.versiones["a1"], usuario=self.revisor, comentario="Propia")
        response = self.client.patch(f"/api/observaciones/{obs.pk}/", {"version": self.versiones["b1"].pk}, format="json")
        self.assertEqual(response.status_code, 400)
        obs.refresh_from_db()
        self.assertEqual(obs.version_id, self.versiones["a1"].pk)
        self.assertEqual(self.client.patch(f"/api/observaciones/{self.obs_ajena.pk}/", {"comentario": "Cambio"}, format="json").status_code, 404)

    def test_creacion_y_actualizacion_asignacion_validan_pareja(self):
        otro = Indicador.objects.create(nombre="Otro", criterio=self.indicador.criterio)
        payload = {"indicador": otro.pk, "unidad_responsable": self.b.pk, "periodo": self.p1.pk}
        self.assertEqual(self.client.post("/api/asignaciones/", payload, format="json").status_code, 400)
        response = self.client.patch(f"/api/asignaciones/{self.asignaciones['a1'].pk}/", {"unidad_responsable": self.b.pk}, format="json")
        self.assertEqual(response.status_code, 400)
        payload["unidad_responsable"] = self.nieta.pk
        self.assertEqual(self.client.post("/api/asignaciones/", payload, format="json").status_code, 201)

    def test_evidencia_existente_no_se_reactiva_fuera_del_ambito(self):
        self.revisor.user_permissions.add(Permission.objects.get(content_type__app_label="evidence", codename="add_evidencia"))
        ajena = self.evidencias["b1"]
        ajena.estado = "cancelada"
        ajena.save()
        response = self.client.post("/api/evidencias/", {"asignacion": ajena.asignacion_id, "titulo": "Ajena", "descripcion": "Ajena"}, format="json")
        self.assertEqual(response.status_code, 400)
        ajena.refresh_from_db()
        self.assertEqual(ajena.estado, "cancelada")

    def test_sin_ambitos_no_hereda_acceso_del_perfil(self):
        self.client.force_authenticate(self.sin_ambitos)
        self.assertEqual(self.ids(self.client.get("/api/asignaciones/")), set())
        self.assertEqual(self.ids(self.client.get("/api/periodos/")), set())
        self.assertEqual(self.client.get("/api/periodos/activo/").status_code, 404)

    def test_periodos_se_limitan_al_alcance(self):
        self.assertEqual(self.ids(self.client.get("/api/periodos/")), {self.p1.pk, self.p2.pk})
        self.assertEqual(self.client.get(f"/api/dashboard/resumen/?periodo={self.p3.pk}").status_code, 404)
        self.assertEqual(self.client.get(f"/api/reportes/general/?periodo={self.p3.pk}").status_code, 404)

    def test_dashboard_y_reportes_no_cruzan_parejas(self):
        for periodo in (self.p1, self.p2):
            response = self.client.get(f"/api/dashboard/resumen/?periodo={periodo.pk}")
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data["asignaciones"], 1)
            response = self.client.get(f"/api/reportes/general/?periodo={periodo.pk}")
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual(response.data["total_asignaciones"], 1)

    def test_reportes_y_exportaciones_excluyen_entregas_ajenas(self):
        Observacion.objects.create(version=self.versiones["a1"], usuario=self.revisor, comentario="Visible")
        response = self.client.get(f"/api/reportes/evidencias/?periodo={self.p1.pk}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual([fila["evidencia"] for fila in response.data["rows"]], ["a1"])
        response = self.client.get(f"/api/reportes/observaciones/?periodo={self.p1.pk}")
        self.assertEqual(response.data["total_observaciones"], 1)
        self.assertEqual(response.data["rows"][0]["comentario"], "Visible")
        self.assertEqual(self.ids(self.client.get(f"/api/observaciones/?periodo={self.p1.pk}")), set(Observacion.objects.filter(comentario="Visible").values_list("pk", flat=True)))
        for recurso in ("evidencias", "observaciones"):
            response = self.client.get(f"/api/reportes/{recurso}/exportar/?periodo={self.p1.pk}&formato=xlsx")
            self.assertEqual(response.status_code, 200)
            libro = load_workbook(BytesIO(response.content), read_only=True)
            celdas = {str(valor) for fila in libro.active.iter_rows(values_only=True) for valor in fila if valor is not None}
            libro.close()
            self.assertIn("a1", celdas)
            self.assertNotIn("b1", celdas)
            self.assertNotIn("Ajena", celdas)

    def test_busqueda_y_correspondencias_conservan_la_pareja_autorizada(self):
        version = VersionMarcoEvaluacion.objects.create(nombre="Marco prueba", slug="revision-prueba")
        dimension = Dimension.objects.create(version_marco=version, codigo="I", nombre="Prueba")
        criterio = Criterio.objects.create(nombre="Referencia exclusiva ajena", periodo=self.p1)
        indicador = Indicador.objects.create(nombre="Indicador exclusivo ajeno", criterio=criterio)
        Asignacion.objects.create(indicador=indicador, unidad_responsable=self.b, periodo=self.p1)
        ajena = Correspondencia.crear_propuesta(actor=self.admin, version_marco=version, dimension=dimension, indicador_op=indicador)
        propia = Correspondencia.crear_propuesta(actor=self.admin, version_marco=version, dimension=dimension, criterio_op=self.indicador.criterio)
        response = self.client.get("/api/vinculacion/correspondencias/")
        self.assertEqual(self.ids(response), {propia.pk})
        self.assertEqual(self.client.get(f"/api/vinculacion/correspondencias/{ajena.pk}/").status_code, 404)
        self.assertEqual(self.client.get(f"/api/vinculacion/contexto/indicador/{indicador.pk}/").status_code, 404)
        response = self.client.get("/api/search/?q=exclusiv")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["indicadores"], [])
        self.assertEqual(response.data["criterios"], [])

    def test_duplicidad_de_ambitos_solapados_no_duplica_asignaciones(self):
        AmbitoRevision.objects.create(usuario=self.revisor, unidad_organizacional=self.a, periodo=self.p1)
        self.assertEqual(self.ids(self.client.get("/api/asignaciones/")), {self.asignaciones["a1"].pk})

    def test_administracion_filtra_y_rechaza_unidad_inactiva(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get(f"/api/ambitos-revision/?usuario={self.revisor.pk}&periodo={self.p1.pk}&activo=true")
        self.assertEqual(self.ids(response), {self.ambito_a.pk})
        self.nieta.activa = False
        self.nieta.save()
        response = self.client.post("/api/ambitos-revision/", {"usuario": self.revisor.pk, "unidad_organizacional": self.nieta.pk, "periodo": self.p1.pk}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_revocacion_invalida_acceso_y_metricas_cacheadas(self):
        url = f"/api/dashboard/resumen/?periodo={self.p1.pk}"
        self.assertEqual(self.client.get(url).data["asignaciones"], 1)
        self.ambito_a.activo = False
        self.ambito_a.save(update_fields=["activo"])
        # P1 ya no está autorizado: no se devuelve el resumen cacheado.
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.ids(self.client.get(f"/api/asignaciones/?periodo={self.p1.pk}")), set())
        self.ambito_b.activo = False
        self.ambito_b.save(update_fields=["activo"])
        self.assertEqual(self.client.get("/api/dashboard/resumen/").data["asignaciones"], 0)

    def test_cambio_de_raiz_no_reutiliza_cache_anterior(self):
        url = f"/api/dashboard/resumen/?periodo={self.p1.pk}"
        self.assertEqual(self.client.get(url).data["asignaciones"], 1)
        self.ambito_a.unidad_organizacional = self.nieta
        self.ambito_a.save()
        self.assertEqual(self.client.get(url).data["asignaciones"], 0)

    def test_ambito_externo_no_autoriza_revision_interna(self):
        AmbitoEvaluacion.objects.create(usuario=self.sin_ambitos, unidad_organizacional=self.raiz_a, periodo=self.p1)
        self.client.force_authenticate(self.sin_ambitos)
        self.assertEqual(self.ids(self.client.get("/api/asignaciones/")), set())

    def test_externo_conserva_restricciones_y_no_gestiona_ambitos_internos(self):
        asig = self.asignaciones["a1"]
        asig.estado = "aprobado"
        asig.save()
        Observacion.objects.create(version=self.versiones["a1"], usuario=self.revisor, comentario="Interna")
        self.client.force_authenticate(self.externo)
        self.assertEqual(self.ids(self.client.get("/api/asignaciones/")), {asig.pk})
        response = self.client.get(f"/api/evidencias/{self.evidencias['a1'].pk}/detalle/")
        self.assertEqual(response.data["historial_estados"], [])
        self.assertEqual(response.data["versiones"][0]["observaciones"], [])
        self.assertEqual(self.client.get("/api/ambitos-revision/").status_code, 403)
        self.assertEqual(self.client.get("/api/ambitos-revision/configuracion/").status_code, 403)
        self.assertEqual(self.client.post(f"/api/asignaciones/{asig.pk}/aprobar/", {}, format="json").status_code, 403)

    def test_administracion_separada_auditada_y_sin_autoconcesion(self):
        payload = {"usuario": self.sin_ambitos.pk, "unidad_organizacional": self.raiz_a.pk, "periodo": self.p1.pk}
        self.assertEqual(self.client.post("/api/ambitos-revision/", payload, format="json").status_code, 403)
        self.client.force_authenticate(self.coordinador)
        response = self.client.post("/api/ambitos-revision/", payload, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        pk = response.data["id"]
        self.assertEqual(self.client.post("/api/ambitos-revision/", payload, format="json").status_code, 400)
        self.assertEqual(self.client.delete(f"/api/ambitos-revision/{pk}/").status_code, 204)
        self.assertFalse(AmbitoRevision.objects.get(pk=pk).activo)
        self.assertTrue(Auditoria.objects.filter(modelo="AmbitoRevision", registro_id=pk, accion="Desactivar ámbito de revisión").exists())
        self.assertEqual(self.client.patch(f"/api/ambitos-revision/{pk}/", {"activo": True}, format="json").status_code, 200)

    def test_no_asignar_ambito_interno_a_usuario_sin_rol(self):
        self.client.force_authenticate(self.admin)
        response = self.client.post("/api/ambitos-revision/", {"usuario": self.responsable.pk, "unidad_organizacional": self.a.pk, "periodo": self.p1.pk}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_catalogo_interno_sigue_siendo_referencia_global(self):
        criterio = Criterio.objects.create(nombre="Referencia no asignada", periodo=self.p3)
        self.assertIn(criterio.pk, self.ids(self.client.get("/api/criterios/")))

    @override_settings(AMBITOS_REVISION_EXPLICITOS=False)
    def test_transicion_conserva_alcance_anterior_hasta_activacion(self):
        self.client.force_authenticate(self.sin_ambitos)
        self.assertEqual(self.ids(self.client.get("/api/asignaciones/")), {self.asignaciones["a1"].pk})
        self.client.force_authenticate(self.admin)
        self.assertEqual(self.client.get("/api/ambitos-revision/configuracion/").data, {"alcance_explicito_activo": False})
