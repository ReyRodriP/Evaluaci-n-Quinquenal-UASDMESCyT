"""Observaciones de evidencias anteriores y decisiones documentales con comentario."""

import io
from collections import Counter
from xml.sax.saxutils import escape

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework.exceptions import PermissionDenied, ValidationError

from accounts.permissions import ROLES_SIN_RESTRICCION, _grupos_usuario, es_evaluador_externo, filtrar_por_rol
from evaluation.filtros import filtrar_por_marco
from evaluation.models import Asignacion
from evaluation.services import relaciones_indicador
from evidence.models import Observacion, RevisionDocumento


def filas_observaciones(request):
    from .views import _formato_fecha, _periodo_para_reporte

    if es_evaluador_externo(request) and not (request.user.is_superuser or _grupos_usuario(request.user) & ROLES_SIN_RESTRICCION):
        raise PermissionDenied('Las observaciones de revisión son de consulta interna.')
    origen = request.query_params.get('origen', '')
    decision = request.query_params.get('decision', '')
    if origen not in ('', 'evidencia', 'documento'):
        raise ValidationError({'origen': 'Seleccione evidencia o documento.'})
    if decision not in ('', 'comentario', 'aprobada', 'observada'):
        raise ValidationError({'decision': 'Seleccione comentario, aprobada u observada.'})
    for campo in ('periodo', 'unidad_organizacional', 'departamento', 'usuario'):
        valor = request.query_params.get(campo)
        if valor and (not str(valor).isdigit() or len(str(valor)) > 19 or not 1 <= int(valor) <= 9223372036854775807):
            raise ValidationError({campo: 'Debe ser un identificador entero positivo.'})
    periodo = _periodo_para_reporte(request)
    asignaciones = filtrar_por_rol(Asignacion.objects.all(), request, dept_field='unidad_responsable')
    asignaciones = asignaciones.filter(periodo=periodo) if periodo else asignaciones.none()
    for campo, lookup in (('unidad_organizacional', 'unidad_responsable_id'), ('departamento', 'unidad_responsable__departamento_legacy_id')):
        valor = request.query_params.get(campo)
        if valor:
            asignaciones = asignaciones.filter(**{lookup: valor})
    asignaciones = filtrar_por_marco(asignaciones, request.query_params)
    usuario = request.query_params.get('usuario')
    filas = []

    def agregar(tipo, recurso, version, observacion, fecha, texto_decision, etiqueta):
        asignacion = recurso.asignacion
        filas.append({
            'evidencia': recurso.titulo, 'indicador': asignacion.indicador_nombre,
            'departamento': asignacion.unidad_responsable.nombre, 'periodo': asignacion.periodo.nombre,
            'unidad_responsable': asignacion.unidad_responsable.nombre,
            'unidad_responsable_id': asignacion.unidad_responsable_id,
            'version': version.numero if tipo == 'documento' else version.version,
            'observador': observacion.usuario.username if observacion.usuario_id else '',
            'observador_id': observacion.usuario_id,
            'comentario': observacion.comentario, 'fecha': _formato_fecha(fecha),
            'tipo_recurso': tipo, 'origen': 'Documento solicitado' if tipo == 'documento' else 'Evidencia anterior',
            'recurso_id': recurso.pk, 'asignacion_id': asignacion.pk, 'periodo_id': asignacion.periodo_id,
            'version_id': version.pk, 'observacion_id': observacion.pk,
            'decision': texto_decision, 'decision_display': etiqueta, '_fecha': fecha,
        })

    if origen != 'documento' and decision not in ('aprobada', 'observada'):
        anteriores = Observacion.objects.filter(activo=True, version__evidencia__asignacion__in=asignaciones)
        anteriores = anteriores.select_related('usuario', 'version__evidencia__asignacion__periodo',
            'version__evidencia__asignacion__unidad_responsable', *relaciones_indicador('version__evidencia__asignacion__'))
        for observacion in anteriores:
            agregar('evidencia', observacion.version.evidencia, observacion.version, observacion,
                observacion.fecha_creacion, 'comentario', 'Comentario')
    if origen != 'evidencia' and decision != 'comentario':
        revisiones = RevisionDocumento.objects.filter(version__solicitud__asignacion__in=asignaciones).exclude(comentario='')
        if decision:
            revisiones = revisiones.filter(decision=decision)
        revisiones = revisiones.select_related('usuario', 'version__solicitud__asignacion__periodo',
            'version__solicitud__asignacion__unidad_responsable', *relaciones_indicador('version__solicitud__asignacion__'))
        for revision in revisiones:
            agregar('documento', revision.version.solicitud, revision.version, revision, revision.fecha,
                revision.decision, 'Aprobada con comentario' if revision.decision == 'aprobada' else 'Cambios solicitados')

    observadores = {f['observador_id']: {'id': f['observador_id'], 'username': f['observador']} for f in filas if f['observador_id']}
    if usuario:
        filas = [f for f in filas if f['observador_id'] == int(usuario)]
    conteos = Counter((f['tipo_recurso'], f['recurso_id']) for f in filas)
    filas.sort(key=lambda f: (f['_fecha'], f['tipo_recurso'], f['observacion_id']), reverse=True)
    for fila in filas:
        fila['total_observaciones'] = conteos[(fila['tipo_recurso'], fila['recurso_id'])]
        del fila['_fecha']
    return filas, len(conteos), sorted(observadores.values(), key=lambda u: (u['username'], u['id']))


def exportar_observaciones(request, filas):
    from .views import COLUMNAS_OBSERVACIONES, _build_pdf

    formato = request.query_params.get('formato', 'xlsx').lower()
    if formato not in ('xlsx', 'pdf'):
        raise ValidationError({'formato': 'Seleccione xlsx o pdf.'})
    columnas = [*COLUMNAS_OBSERVACIONES, 'Origen', 'Decisión', 'ID asignación', 'ID recurso', 'ID versión']
    campos = ['evidencia', 'indicador', 'departamento', 'periodo', 'version', 'observador', 'comentario',
        'fecha', 'total_observaciones', 'origen', 'decision_display', 'asignacion_id', 'recurso_id', 'version_id']
    valores = [[fila[campo] for campo in campos] for fila in filas]
    if formato == 'pdf':
        # La trazabilidad numérica completa se conserva en Excel; PDF prioriza el texto de la observación.
        seleccion = [0, 1, 2, 3, 4, 5, 6, 7, 9, 10]
        contenido = _build_pdf('Reporte de Observaciones', [columnas[i] for i in seleccion],
            [[escape(str(fila[i])) for i in seleccion] for fila in valores], split_in_row=True)
        tipo = 'application/pdf'
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = 'Observaciones'
        hoja.append(columnas)
        for celda in hoja[1]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='1F4E79')
        for fila in valores:
            hoja.append(fila)
            for celda in hoja[hoja.max_row]:
                if isinstance(celda.value, str):
                    celda.data_type = 's'
                celda.alignment = Alignment(wrap_text=True, vertical='top')
        for columna in hoja.columns:
            hoja.column_dimensions[columna[0].column_letter].width = 30 if columna[0].column == 7 else 22
        hoja.freeze_panes = 'A2'
        buffer = io.BytesIO()
        libro.save(buffer)
        contenido = buffer.getvalue()
        tipo = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    respuesta = HttpResponse(contenido, content_type=tipo)
    respuesta['Content-Disposition'] = f'attachment; filename="reporte_observaciones.{formato}"'
    return respuesta
