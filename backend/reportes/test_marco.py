"""Resultados del marco: conteos distintos, cierre explícito y alcance de exportación."""
import csv
import io

from django.test import TestCase, override_settings
from openpyxl import load_workbook

from evaluation.models import Asignacion
from evidence.test_documentos import EscenarioDocumental
from organization.models import AmbitoRevision


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ReporteMarcoTests(EscenarioDocumental, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.crear_escenario()

    def setUp(self):
        self.preparar()

    def consultar(self, usuario=None, **params):
        response = self.cliente(usuario).get('/api/reportes/marco/', {'periodo': self.periodo.pk, **params})
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def test_documentos_y_versiones_no_duplican_asignaciones(self):
        a, b = self.solicitar(), self.solicitar('Documento B')
        self.aprobar_documento(a)
        self.assertEqual(self.subir(b).status_code, 201)
        self.assertEqual(self.subir(b).status_code, 201)
        self.enviar(b)
        data = self.consultar(agrupar='estandar')
        self.assertEqual(data['total'], 1)
        resumen = data['resumen']
        self.assertEqual(resumen['asignaciones'], 1)
        self.assertEqual(resumen['con_archivos'], 1)
        self.assertEqual(resumen['documentos_solicitados'], 2)
        self.assertEqual(resumen['documentos_entregados'], 2)
        self.assertEqual(resumen['documentos_aprobados'], 1)
        self.assertEqual(resumen['documentos_en_revision'], 1)
        self.assertEqual(resumen['aprobadas_globalmente'], 0)
        self.assertEqual(data['rows'][0]['estandar']['codigo'], self.marco.estandar.codigo_completo)
        self.assertFalse(self.version.activo)
        self.assertEqual(data['version_periodo']['id'], self.version.pk)

    def test_documentos_aprobados_requieren_cierre_global(self):
        a, b = self.solicitar(), self.solicitar('Documento B')
        for pk in (a, b):
            self.aprobar_documento(pk)
        resumen = self.consultar()['resumen']
        self.assertEqual(resumen['pendientes_cierre'], 1)
        self.assertEqual(resumen['aprobadas_globalmente'], 0)
        self.assertEqual(resumen['aprobacion_global_porcentaje'], 0)
        self.assertEqual(self.cliente(self.revisor).post(f'/api/asignaciones/{self.asignacion.pk}/aprobar/').status_code, 200)
        resumen = self.consultar()['resumen']
        self.assertEqual(resumen['pendientes_cierre'], 0)
        self.assertEqual(resumen['aprobadas_globalmente'], 1)
        self.assertEqual(resumen['aprobacion_global_porcentaje'], 100)

    def test_observacion_y_borrador_se_distinguen_de_revision(self):
        pk = self.solicitar()
        self.assertEqual(self.subir(pk).status_code, 201)
        self.enviar(pk)
        self.assertEqual(self.revisar(pk, 'solicitar_cambios', 'Corregir contenido').status_code, 200)
        resumen = self.consultar()['resumen']
        self.assertEqual(resumen['documentos_observados'], 1)
        self.assertEqual(resumen['documentos_en_revision'], 0)
        self.assertEqual(self.subir(pk).status_code, 201)
        resumen = self.consultar()['resumen']
        self.assertEqual(resumen['documentos_observados'], 0)
        self.assertEqual(resumen['documentos_en_revision'], 0)
        self.assertEqual(resumen['documentos_entregados'], 1)

    def test_revisor_solo_ve_asignaciones_y_opciones_autorizadas(self):
        Asignacion.objects.create(indicador_marco=self.marco_otra_dimension, periodo=self.periodo,
            unidad_responsable=self.ajena, encargado=self.otro)
        self.assertEqual(self.consultar()['resumen']['asignaciones'], 2)
        data = self.consultar(self.revisor)
        self.assertEqual(data['resumen']['asignaciones'], 1)
        self.assertEqual([d['id'] for d in data['opciones']['dimensiones']], [self.marco.estandar.criterio_marco.componente.dimension_id])
        self.assertEqual(self.consultar(self.revisor, unidad=self.ajena.pk)['resumen']['asignaciones'], 0)
        self.ambito.activo = False
        self.ambito.save()
        self.assertEqual(self.cliente(self.revisor).get('/api/reportes/marco/', {'periodo': self.periodo.pk}).status_code, 404)
        self.assertEqual(self.cliente(self.revisor).get('/api/reportes/marco/exportar/', {'periodo': self.periodo.pk, 'formato': 'csv'}).status_code, 404)

    def test_periodos_historicos_no_cruzan_ambitos(self):
        Asignacion.objects.create(indicador_marco=self.marco, periodo=self.historico,
            unidad_responsable=self.unidad, encargado=self.responsable)
        AmbitoRevision.objects.create(usuario=self.revisor, periodo=self.historico, unidad_organizacional=self.ajena)
        data = self.consultar(self.revisor, periodo=self.historico.pk)
        self.assertEqual(data['resumen']['asignaciones'], 0)
        self.assertEqual(self.consultar(periodo=self.historico.pk)['resumen']['asignaciones'], 1)

    def test_filtros_y_parametros_invalidos(self):
        Asignacion.objects.create(indicador_marco=self.marco_otra_dimension, periodo=self.periodo,
            unidad_responsable=self.unidad, encargado=self.responsable)
        self.assertEqual(self.consultar()['total'], 2)
        self.assertEqual(self.consultar(estandar=self.marco.estandar_id)['resumen']['asignaciones'], 1)
        for params in ({'agrupar': 'invalido'}, {'dimension': 'texto'}, {'periodo': 'texto'}, {'unidad': '-1'}):
            response = self.cliente().get('/api/reportes/marco/', {'periodo': self.periodo.pk, **params})
            self.assertEqual(response.status_code, 400, response.data)
        for usuario in (self.responsable, self.externo):
            self.assertEqual(self.cliente(usuario).get('/api/reportes/marco/').status_code, 403)
            self.assertEqual(self.cliente(usuario).get('/api/reportes/marco/exportar/').status_code, 403)

    def test_exporta_todos_los_grupos_filtrados_y_porcentajes_numericos(self):
        Asignacion.objects.create(indicador_marco=self.marco_otra_dimension, periodo=self.periodo,
            unidad_responsable=self.unidad, encargado=self.responsable)
        data = self.consultar(page_size=1)
        self.assertEqual(len(data['rows']), 1)
        self.assertEqual(data['total'], 2)
        params = {'periodo': self.periodo.pk, 'page_size': 1}
        response = self.cliente().get('/api/reportes/marco/exportar/', {**params, 'formato': 'csv'})
        self.assertEqual(response.status_code, 200)
        filas = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(filas), 3)
        self.assertTrue(all(len(fila) == 14 for fila in filas))
        response = self.cliente().get('/api/reportes/marco/exportar/', {**params, 'formato': 'xlsx'})
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(hoja.max_row, 3)
        self.assertEqual(hoja['N2'].data_type, 'n')
        self.assertEqual(hoja['N2'].value, 0)
        response = self.cliente().get('/api/reportes/marco/exportar/', {**params, 'formato': 'pdf'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_exportacion_filtra_ambito_y_texto_no_se_ejecuta_como_formula(self):
        self.periodo.nombre = '=1+1'
        self.periodo.save()
        Asignacion.objects.create(indicador_marco=self.marco_otra_dimension, periodo=self.periodo,
            unidad_responsable=self.ajena, encargado=self.otro)
        params = {'periodo': self.periodo.pk}
        response = self.cliente(self.revisor).get('/api/reportes/marco/exportar/', {**params, 'formato': 'xlsx'})
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(hoja['A2'].data_type, 's')
        self.assertEqual(hoja['A2'].value, '=1+1')
        response = self.cliente(self.revisor).get('/api/reportes/marco/exportar/', {**params, 'formato': 'csv'})
        filas = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[1][0], "'=1+1")
