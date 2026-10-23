"""Comentarios históricos por recurso, versión y ámbito, sin inferir decisiones antiguas."""
import io

from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from evaluation.models import Asignacion
from evidence.models import (
    Evidencia,
    Observacion,
    RevisionDocumento,
    SolicitudDocumental,
    VersionDocumento,
    VersionEvidencia,
)
from evidence.test_documentos import EscenarioDocumental
from organization.models import AmbitoRevision


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class ObservacionesDocumentalesTests(EscenarioDocumental, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.crear_escenario()

    def setUp(self):
        self.preparar()

    def consultar(self, actor=None, **filtros):
        response = self.cliente(actor).get('/api/reportes/observaciones/', {'periodo': self.periodo.pk, **filtros})
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def observar(self, titulo='Documento', comentario='Corregir la fecha'):
        pk = self.solicitar(titulo)
        self.assertEqual(self.subir(pk).status_code, 201)
        self.enviar(pk)
        self.assertEqual(self.revisar(pk, 'solicitar_cambios', comentario).status_code, 200)
        return pk

    def comentario_anterior(self, titulo='Documento', pk=None):
        evidencia = Evidencia.objects.create(asignacion=self.legada, titulo=titulo, descripcion='Anterior', **({'id_evidencia': pk} if pk else {}))
        version = VersionEvidencia.objects.create(evidencia=evidencia, version=1, archivo=SimpleUploadedFile('anterior.pdf', b'%PDF-1.4\nprueba'))
        return Observacion.objects.create(version=version, usuario=self.revisor, comentario='Comentario anterior')

    def test_recursos_con_mismo_titulo_e_id_se_cuentan_por_origen(self):
        pk = self.observar()
        anterior = self.comentario_anterior(pk=pk)
        data = self.consultar()
        self.assertEqual(data['total_observaciones'], 2)
        self.assertEqual(data['evidencias_observadas'], 2)
        self.assertEqual({r['tipo_recurso'] for r in data['rows']}, {'documento', 'evidencia'})
        doc = next(r for r in data['rows'] if r['tipo_recurso'] == 'documento')
        self.assertEqual(doc['recurso_id'], pk)
        self.assertEqual(doc['version_id'], VersionDocumento.objects.get(solicitud_id=pk).pk)
        self.assertEqual(doc['decision'], 'observada')
        old = next(r for r in data['rows'] if r['tipo_recurso'] == 'evidencia')
        self.assertEqual(old['version_id'], anterior.version_id)
        self.assertEqual(old['decision'], 'comentario')
        self.assertTrue(all(r['total_observaciones'] == 1 for r in data['rows']))

    def test_correccion_conserva_comentario_y_version_anteriores(self):
        pk = self.observar(comentario='Observación de versión 1')
        version_1 = VersionDocumento.objects.get(solicitud_id=pk)
        self.assertEqual(self.subir(pk).status_code, 201)
        self.enviar(pk)
        self.assertEqual(self.revisar(pk, comentario='Versión 2 validada').status_code, 200)
        data = self.consultar(origen='documento')
        self.assertEqual(data['total_observaciones'], 2)
        self.assertEqual(data['evidencias_observadas'], 1)
        self.assertEqual([r['version'] for r in data['rows']], [2, 1])
        self.assertTrue(all(r['total_observaciones'] == 2 for r in data['rows']))
        old = next(r for r in data['rows'] if r['comentario'] == 'Observación de versión 1')
        self.assertEqual(old['version_id'], version_1.pk)
        self.assertEqual(old['decision'], 'observada')

    def test_aprobacion_sin_comentario_no_inventa_observaciones(self):
        pk = self.observar()
        self.assertEqual(self.subir(pk).status_code, 201)
        self.enviar(pk)
        self.assertEqual(self.revisar(pk).status_code, 200)
        self.assertEqual(self.consultar()['total_observaciones'], 1)
        self.assertEqual(self.consultar(decision='aprobada')['total_observaciones'], 0)

    def test_filtros_por_origen_decision_usuario_y_observaciones_activas(self):
        self.observar()
        anterior = self.comentario_anterior()
        self.assertEqual(self.consultar(origen='documento', decision='observada')['total_observaciones'], 1)
        self.assertEqual(self.consultar(origen='evidencia', decision='aprobada')['total_observaciones'], 0)
        self.assertEqual(self.consultar(usuario=self.coordinador.pk)['total_observaciones'], 0)
        self.assertEqual(self.consultar(usuario=self.revisor.pk)['total_observaciones'], 2)
        anterior.activo = False
        anterior.save()
        self.assertEqual(self.consultar(origen='evidencia')['total_observaciones'], 0)
        for filtros in ({'origen': 'otro'}, {'decision': 'rechazada'}, {'usuario': '-1'}, {'periodo': 'texto'}):
            self.assertEqual(self.cliente().get('/api/reportes/observaciones/', {'periodo': self.periodo.pk, **filtros}).status_code, 400)

    def test_ambito_limita_comentarios_observadores_y_exportacion(self):
        self.observar()
        ajena = Asignacion.objects.create(indicador_marco=self.marco, periodo=self.periodo, unidad_responsable=self.ajena, encargado=self.otro, gestion_documental=True)
        solicitud = SolicitudDocumental.objects.create(asignacion=ajena, titulo='Ajena', descripcion='Fuera de ámbito', creada_por=self.coordinador)
        version = VersionDocumento.objects.create(solicitud=solicitud, numero=1, archivo=SimpleUploadedFile('ajena.pdf', b'%PDF'))
        RevisionDocumento.objects.create(version=version, usuario=self.coordinador, decision='observada', comentario='Comentario ajeno')
        self.assertEqual(self.consultar()['total_observaciones'], 2)
        data = self.consultar(self.revisor)
        self.assertEqual(data['total_observaciones'], 1)
        self.assertEqual(data['observadores'], [{'id': self.revisor.pk, 'username': self.revisor.username}])
        response = self.cliente(self.revisor).get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk, 'formato': 'xlsx'})
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(hoja.max_row, 2)
        self.ambito.activo = False
        self.ambito.save()
        for endpoint in ('/api/reportes/observaciones/', '/api/reportes/observaciones/exportar/'):
            self.assertEqual(self.cliente(self.revisor).get(endpoint, {'periodo': self.periodo.pk}).status_code, 404)

    def test_historia_y_ambitos_no_se_cruzan_entre_periodos(self):
        self.observar()
        self.periodo.activo = False
        self.periodo.save()
        self.assertEqual(self.consultar(self.revisor)['total_observaciones'], 1)
        historica = Asignacion.objects.create(indicador_marco=self.marco, periodo=self.historico, unidad_responsable=self.unidad, encargado=self.responsable, gestion_documental=True)
        solicitud = SolicitudDocumental.objects.create(asignacion=historica, titulo='Histórica', descripcion='Conservar', creada_por=self.coordinador)
        version = VersionDocumento.objects.create(solicitud=solicitud, numero=1, archivo=SimpleUploadedFile('historica.pdf', b'%PDF'))
        RevisionDocumento.objects.create(version=version, usuario=self.revisor, decision='observada', comentario='Comentario histórico')
        AmbitoRevision.objects.create(usuario=self.revisor, periodo=self.historico, unidad_organizacional=self.ajena)
        self.assertEqual(self.consultar(periodo=self.historico.pk)['total_observaciones'], 1)
        self.assertEqual(self.consultar(self.revisor, periodo=self.historico.pk)['total_observaciones'], 0)

    def test_exportacion_conserva_filtros_ids_y_texto_literal(self):
        self.observar(comentario='=SUM(1,1) <nombre> & "fecha"')
        self.comentario_anterior()
        response = self.cliente().get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk, 'origen': 'documento', 'formato': 'xlsx', 'page_size': 1})
        hoja = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(hoja.max_row, 2)
        self.assertEqual(hoja['G2'].value, '=SUM(1,1) <nombre> & "fecha"')
        self.assertEqual(hoja['G2'].data_type, 's')
        self.assertEqual(hoja['J2'].value, 'Documento solicitado')
        self.assertEqual(hoja['N2'].value, VersionDocumento.objects.get().pk)
        response = self.cliente().get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk, 'formato': 'pdf'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
        response = self.cliente().get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk, 'formato': 'xlsx', 'page_size': 1})
        self.assertEqual(load_workbook(io.BytesIO(response.content)).active.max_row, 3)

    def test_observaciones_internas_no_se_conceden_a_roles_externos(self):
        self.observar()
        for usuario in (self.externo, self.responsable):
            self.assertEqual(self.cliente(usuario).get('/api/reportes/observaciones/').status_code, 403)
        self.revisor.groups.add(Group.objects.get(name='Evaluador Externo'))
        self.assertEqual(self.cliente(self.revisor).get('/api/reportes/observaciones/', {'periodo': self.periodo.pk}).status_code, 403)
        self.assertEqual(self.cliente(self.revisor).get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk}).status_code, 403)

    def test_pdf_admite_comentarios_que_necesitan_varias_paginas(self):
        self.observar(comentario='Revisar <fecha> & documento. ' * 300)
        response = self.cliente().get('/api/reportes/observaciones/exportar/', {'periodo': self.periodo.pk, 'formato': 'pdf'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b'%PDF'))
