"""Escenario sintético de 3B: referencias normativas, compatibilidad y entregas."""

import csv
import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from threading import Event
from unittest.mock import patch

from django.contrib.auth.models import Group
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, close_old_connections, connection, transaction
from django.test import TestCase, TransactionTestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import Usuario
from auditoria.models import Auditoria
from catalogo.models import Componente, CriterioMarco, Dimension, Estandar, IndicadorMarco, VersionMarcoEvaluacion
from evaluation.models import Asignacion, Criterio, HistorialEstado, Indicador, Periodo
from evidence.models import Evidencia, Observacion, VersionEvidencia
from notificaciones.models import Notificacion
from organization.models import (
    AmbitoEvaluacion,
    AmbitoRevision,
    PerfilUsuario,
    TipoUnidadOrganizacional,
    UnidadOrganizacional,
)
from vinculacion.models import Correspondencia


class EscenarioMarco:
    @classmethod
    def crear_escenario(cls):
        def usuario(nombre, rol):
            u = Usuario.objects.create_user(username=nombre, email=f"{nombre}@example.test", password="testpass")
            u.groups.add(Group.objects.get_or_create(name=rol)[0])
            return u

        cls.coordinador = usuario("coord_3b", "Coordinador Quinquenal")
        cls.responsable = usuario("resp_3b", "Responsable Departamental")
        cls.colaborador = usuario("colab_3b", "Responsable Departamental")
        cls.revisor = usuario("rev_3b", "Revisor Institucional")
        cls.externo = usuario("ext_3b", "Evaluador Externo")
        cls.otro = usuario("otro_3b", "Responsable Departamental")
        tipo = TipoUnidadOrganizacional.objects.get_or_create(nombre="Centro")[0]
        cls.raiz = UnidadOrganizacional.objects.create(nombre="Centro técnico 3B", tipo=tipo)
        cls.unidad = UnidadOrganizacional.objects.create(nombre="Unidad técnica 3B", tipo=tipo, unidad_padre=cls.raiz)
        cls.ajena = UnidadOrganizacional.objects.create(nombre="Unidad ajena 3B", tipo=tipo)
        for u in (cls.responsable, cls.colaborador):
            PerfilUsuario.objects.create(usuario=u, unidad_organizacional=cls.unidad)
        PerfilUsuario.objects.create(usuario=cls.otro, unidad_organizacional=cls.ajena)
        cls.version = VersionMarcoEvaluacion.objects.create(nombre="Marco técnico de prueba A", slug="tecnico-3b-a", activo=False)
        cls.version_b = VersionMarcoEvaluacion.objects.create(nombre="Marco técnico de prueba B", slug="tecnico-3b-b", activo=False)

        def rama(version, codigo, texto):
            dimension = Dimension.objects.create(version_marco=version, codigo=codigo, nombre=f"Dimensión técnica {texto}")
            comp = Componente.objects.create(version_marco=version, dimension=dimension, codigo=f"{codigo}.1", codigo_completo=f"{codigo}.1", nombre="Componente técnico")
            criterio = CriterioMarco.objects.create(version_marco=version, componente=comp, codigo=f"{codigo}.1.1", codigo_completo=f"{codigo}.1.1", nombre=f"Criterio técnico {texto}")
            estandar = Estandar.objects.create(version_marco=version, criterio_marco=criterio, codigo=f"{codigo}.1.1.1", codigo_completo=f"{codigo}.1.1.1", redaccion="Estándar sintético de prueba")
            indicador = IndicadorMarco.objects.create(version_marco=version, estandar=estandar, codigo=f"{estandar.codigo}-IND-01", codigo_completo=f"{estandar.codigo}-IND-01", redaccion_oficial=f"Indicador técnico {texto}", orden=1)
            return indicador

        cls.marco = rama(cls.version, "I", "A")
        cls.marco_hermano = IndicadorMarco.objects.create(version_marco=cls.version, estandar=cls.marco.estandar, codigo="I.1.1.1-IND-02", codigo_completo="I.1.1.1-IND-02", redaccion_oficial="Segundo indicador técnico", orden=2)
        cls.marco_otra_dimension = rama(cls.version, "II", "Otra dimensión")
        cls.marco_b = rama(cls.version_b, "I", "B")
        hoy = timezone.localdate()
        cls.periodo = Periodo.objects.create(nombre="Período técnico 3B", fecha_inicio=hoy, fecha_fin=hoy + timedelta(days=365), version_marco=cls.version)
        cls.historico = Periodo.objects.create(nombre="Histórico técnico", fecha_inicio=hoy-timedelta(days=730), fecha_fin=hoy-timedelta(days=365), activo=False, version_marco=cls.version)
        cls.criterio = Criterio.objects.create(nombre="Criterio operativo técnico", periodo=cls.periodo)
        cls.operativo = Indicador.objects.create(pk=cls.marco.pk, nombre="Indicador operativo con ID coincidente", criterio=cls.criterio)
        cls.legada = Asignacion.objects.create(indicador=cls.operativo, unidad_responsable=cls.unidad, periodo=cls.periodo)
        cls.ambito = AmbitoRevision.objects.create(usuario=cls.revisor, unidad_organizacional=cls.raiz, periodo=cls.periodo)
        AmbitoEvaluacion.objects.create(usuario=cls.externo, unidad_organizacional=cls.raiz, periodo=cls.periodo)

    def payload(self, **cambios):
        return {"indicador_marco": self.marco.pk, "unidad_responsable": self.unidad.pk, "periodo": self.periodo.pk, "encargado": self.responsable.pk, **cambios}

    def cliente(self, user=None):
        c = APIClient()
        c.force_authenticate(user or self.coordinador)
        return c


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class AsignacionesMarcoTests(EscenarioMarco, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.crear_escenario()

    def setUp(self):
        cache.clear()

    def crear(self, **cambios):
        response = self.cliente().post("/api/asignaciones/", self.payload(**cambios), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        return Asignacion.objects.get(pk=response.data["id"])

    def test_crear_sin_indicador_operativo_con_ruta_y_encargado(self):
        antes = (Indicador.objects.count(), Criterio.objects.count(), Correspondencia.objects.count())
        response = self.cliente().post("/api/asignaciones/", self.payload(), format="json")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data["indicador"])
        self.assertEqual(response.data["tipo_origen"], "marco")
        self.assertEqual(response.data["indicador_nombre"], self.marco.redaccion_oficial)
        self.assertEqual(response.data["encargado_nombre"], self.responsable.username)
        ruta = response.data["ruta_marco"]
        self.assertEqual(ruta["version"]["id"], self.version.pk)
        self.assertFalse(ruta["version"]["activo"])
        self.assertEqual(ruta["estandar"]["id"], self.marco.estandar_id)
        self.assertTrue(ruta["indicador"]["codigo_sintetico"])
        self.assertEqual((Indicador.objects.count(), Criterio.objects.count(), Correspondencia.objects.count()), antes)
        self.assertTrue(Auditoria.objects.filter(modelo="Asignacion", registro_id=response.data["id"], accion="Crear asignación").exists())
        self.assertTrue(Notificacion.objects.filter(usuario=self.responsable, titulo="Nuevo indicador asignado").exists())

    def test_legado_conserva_contrato_y_campos_vacios(self):
        response = self.cliente().get(f"/api/asignaciones/{self.legada.pk}/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["indicador"], self.operativo.pk)
        self.assertIsNone(response.data["indicador_marco"])
        self.assertIsNone(response.data["ruta_marco"])
        self.assertIsNone(response.data["encargado"])
        self.assertEqual(response.data["indicador_nombre"], self.operativo.nombre)

    def test_rechaza_ambos_origenes_y_ninguno(self):
        for cambios in ({"indicador": self.operativo.pk}, {"indicador_marco": None}):
            r = self.cliente().post("/api/asignaciones/", self.payload(**cambios), format="json")
            self.assertEqual(r.status_code, 400, r.data)

    def test_solo_se_crea_pendiente(self):
        r = self.cliente().post("/api/asignaciones/", self.payload(estado="aprobado"), format="json")
        self.assertEqual(r.status_code, 400)

    def test_version_de_periodo_debe_seleccionarse_explicitamente(self):
        self.periodo.version_marco = None
        self.periodo.save()
        r = self.cliente().post("/api/asignaciones/", self.payload(), format="json")
        self.assertEqual(r.status_code, 400)
        r = self.cliente().patch(f"/api/periodos/{self.periodo.pk}/", {"version_marco": self.version.pk}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.version.refresh_from_db()
        self.assertFalse(self.version.activo)
        self.crear()

    def test_rechaza_version_ajena_y_ancestros_incoherentes(self):
        r = self.cliente().post("/api/asignaciones/", self.payload(indicador_marco=self.marco_b.pk), format="json")
        self.assertEqual(r.status_code, 400)
        incoherente = VersionMarcoEvaluacion.objects.create(nombre="Versión de prueba incoherente", slug="tecnico-incoherente")
        for objeto in (self.marco.estandar, self.marco.estandar.criterio_marco, self.marco.estandar.criterio_marco.componente, self.marco.estandar.criterio_marco.componente.dimension):
            with self.subTest(nivel=type(objeto).__name__):
                type(objeto).objects.filter(pk=objeto.pk).update(version_marco=incoherente)
                r = self.cliente().post("/api/asignaciones/", self.payload(), format="json")
                self.assertEqual(r.status_code, 400, r.data)
                type(objeto).objects.filter(pk=objeto.pk).update(version_marco=self.version)

    def test_encargado_obligatorio_activo_de_unidad_y_con_permisos(self):
        for encargado in (None, self.otro.pk):
            self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(encargado=encargado), format="json").status_code, 400)
        self.responsable.is_active = False
        self.responsable.save()
        self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(), format="json").status_code, 400)
        self.responsable.is_active = True
        self.responsable.save()
        self.responsable.groups.clear()
        self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(), format="json").status_code, 400)

    def test_indicador_inactivo_e_historico_no_admiten_nuevas_asignaciones(self):
        self.marco.activo = False
        self.marco.save()
        self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(), format="json").status_code, 400)
        self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(periodo=self.historico.pk), format="json").status_code, 400)

    def test_no_sustituir_version_con_asignaciones_por_api_ni_modelo(self):
        self.crear()
        r = self.cliente().patch(f"/api/periodos/{self.periodo.pk}/", {"version_marco": self.version_b.pk}, format="json")
        self.assertEqual(r.status_code, 400)
        self.periodo.version_marco = self.version_b
        with self.assertRaises(ValidationError):
            self.periodo.save()
        self.periodo.refresh_from_db()
        self.assertEqual(self.periodo.version_marco_id, self.version.pk)

    def test_unicidad_y_origen_sql(self):
        asignacion = self.crear()
        self.assertEqual(self.cliente().post("/api/asignaciones/", self.payload(), format="json").status_code, 400)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Asignacion.objects.create(indicador_marco=self.marco, unidad_responsable=self.unidad, periodo=self.periodo, encargado=self.responsable)
        for campos in ({"indicador_id": self.operativo.pk}, {"indicador_marco_id": None}):
            with self.subTest(campos=campos), self.assertRaises(IntegrityError), transaction.atomic():
                Asignacion.objects.filter(pk=asignacion.pk).update(**campos)

    def test_modelo_y_formularios_no_admiten_marco_sin_encargado(self):
        asignacion = Asignacion(indicador_marco=self.marco, unidad_responsable=self.unidad, periodo=self.periodo)
        with self.assertRaises(ValidationError):
            asignacion.full_clean()
        with self.assertRaises(ValidationError):
            asignacion.save()

    def test_mismo_indicador_puede_asignarse_a_otra_unidad(self):
        self.crear()
        otro = self.crear(unidad_responsable=self.ajena.pk, encargado=self.otro.pk)
        self.assertEqual(otro.indicador_marco_id, self.marco.pk)

    def test_filtros_de_ruta_y_colision_de_ids_operativos(self):
        asignacion = self.crear()
        self.crear(indicador_marco=self.marco_otra_dimension.pk)
        ruta = asignacion.ruta_marco
        filtros = {"version_marco": self.version.pk, "dimension": ruta["dimension"]["id"], "componente": ruta["componente"]["id"], "criterio_marco": ruta["criterio"]["id"], "estandar": ruta["estandar"]["id"], "indicador_marco": self.marco.pk}
        for campo, valor in filtros.items():
            r = self.cliente().get("/api/asignaciones/", {campo: valor})
            self.assertEqual(r.status_code, 200, r.data)
            ids = {a["id"] for a in r.data["results"]}
            self.assertIn(asignacion.pk, ids)
            self.assertNotIn(self.legada.pk, ids)
            if campo != "version_marco":
                self.assertEqual(ids, {asignacion.pk})
        r = self.cliente().get("/api/asignaciones/", {"indicador": self.marco.pk})
        self.assertEqual({a["id"] for a in r.data["results"]}, {self.legada.pk})
        self.assertEqual(self.cliente().get("/api/asignaciones/", {"dimension": "invalida"}).status_code, 400)
        self.assertEqual(self.cliente().get("/api/asignaciones/", {"dimension": ruta["dimension"]["id"], "estandar": self.marco_otra_dimension.estandar_id}).data["count"], 0)

    def test_filtros_no_amplian_alcance_revisor(self):
        asignacion = self.crear()
        self.crear(unidad_responsable=self.ajena.pk, encargado=self.otro.pk)
        r = self.cliente(self.revisor).get("/api/asignaciones/", {"indicador_marco": self.marco.pk})
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["results"][0]["id"], asignacion.pk)
        self.ambito.activo = False
        self.ambito.save()
        self.assertEqual(self.cliente(self.revisor).get("/api/asignaciones/", {"indicador_marco": self.marco.pk}).data["count"], 0)

    def test_patch_parcial_no_pierde_origen_encargado_y_no_reclasifica_documentos(self):
        asignacion = self.crear()
        r = self.cliente().patch(f"/api/asignaciones/{asignacion.pk}/", {"encargado": self.colaborador.pk}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["indicador_marco"], self.marco.pk)
        Evidencia.objects.create(asignacion=asignacion, titulo="Documento", descripcion="Prueba")
        r = self.cliente().patch(f"/api/asignaciones/{asignacion.pk}/", {"indicador_marco": self.marco_hermano.pk}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_fallo_de_auditoria_revierte_la_creacion(self):
        with patch("evaluation.views.registrar_auditoria", side_effect=RuntimeError("fallo simulado")), self.assertRaises(RuntimeError):
            self.cliente().post("/api/asignaciones/", self.payload(), format="json")
        self.assertFalse(Asignacion.objects.filter(indicador_marco__isnull=False).exists())

    def test_consulta_historica_con_marco_conserva_la_ruta(self):
        asignacion = Asignacion.objects.create(indicador_marco=self.marco, unidad_responsable=self.unidad, periodo=self.historico, encargado=self.responsable)
        AmbitoRevision.objects.create(usuario=self.revisor, unidad_organizacional=self.raiz, periodo=self.historico)
        c = self.cliente(self.revisor)
        response = c.get(f"/api/asignaciones/{asignacion.pk}/", {"periodo": self.historico.pk})
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["ruta_marco"]["indicador"]["id"], self.marco.pk)
        response = c.patch(f"/api/asignaciones/{asignacion.pk}/", {"estado": "en_progreso"}, format="json")
        self.assertEqual(response.status_code, 400)

    def test_flujo_documental_colaborativo_y_revision_con_auditoria(self):
        asignacion = self.crear()
        responsable = self.cliente(self.responsable)
        r = responsable.post("/api/evidencias/", {"asignacion": asignacion.pk, "titulo": "Entrega técnica", "descripcion": "Prueba 3B"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        evidencia = Evidencia.objects.get(pk=r.data["id_evidencia"])
        r = responsable.post(f"/api/evidencias/{evidencia.pk}/subir_version/", {"archivo": SimpleUploadedFile("v1.txt", b"Version 1", content_type="text/plain")}, format="multipart")
        self.assertEqual(r.status_code, 201, r.data)
        version1 = VersionEvidencia.objects.get(pk=r.data["id_version"])
        revisor = self.cliente(self.revisor)
        r = revisor.post(f"/api/asignaciones/{asignacion.pk}/observada/", {"comentario": "Corregir entrega"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertTrue(Observacion.objects.filter(version=version1, comentario="Corregir entrega").exists())
        r = self.cliente(self.colaborador).post(f"/api/evidencias/{evidencia.pk}/subir_version/", {"archivo": SimpleUploadedFile("v2.txt", b"Version 2", content_type="text/plain")}, format="multipart")
        self.assertEqual(r.status_code, 201, r.data)
        r = revisor.post(f"/api/asignaciones/{asignacion.pk}/aprobar/", {}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(VersionEvidencia.objects.filter(evidencia=evidencia).count(), 2)
        self.assertEqual(HistorialEstado.objects.filter(asignacion=asignacion).count(), 4)
        self.assertTrue(Auditoria.objects.filter(modelo="Asignacion", registro_id=asignacion.pk, accion="Aprobar evidencia").exists())
        r = revisor.get(f"/api/evidencias/{evidencia.pk}/detalle/")
        self.assertEqual(r.data["asignacion_info"]["ruta_marco"]["indicador"]["id"], self.marco.pk)
        r = self.cliente(self.externo).get(f"/api/evidencias/{evidencia.pk}/detalle/")
        self.assertEqual(len(r.data["versiones"]), 1)
        self.assertEqual(r.data["versiones"][0]["observaciones"], [])
        self.assertEqual(r.data["historial_estados"], [])

    def test_evidencias_zip_reportes_y_metricas_con_dos_origenes(self):
        asignacion = self.crear()
        evidencia = Evidencia.objects.create(asignacion=asignacion, titulo="Documento marco", descripcion="Prueba")
        VersionEvidencia.objects.create(evidencia=evidencia, archivo=SimpleUploadedFile("marco.txt", b"contenido", content_type="text/plain"), version=1)
        c = self.cliente()
        r = c.get("/api/evidencias/", {"dimension": self.marco.estandar.criterio_marco.componente.dimension_id})
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["count"], 1)
        r = c.post("/api/evidencias/descargas-masivas/", {"evidencias": [evidencia.pk]}, format="json")
        self.assertEqual(r.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content))) as z:
            manifest = list(csv.DictReader(io.StringIO(z.read("manifest.csv").decode("utf-8-sig"))))
            self.assertEqual(manifest[0]["indicador"], self.marco.redaccion_oficial)
            self.assertEqual(manifest[0]["criterio"], self.marco.estandar.criterio_marco.nombre)
        r = c.get("/api/dashboard/resumen/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["indicadores"], 2)
        self.assertEqual(r.data["indicadores_operativos"], 1)
        self.assertEqual(r.data["indicadores_marco"], 1)
        self.assertEqual(c.get(f"/api/dashboard/periodo/{self.periodo.pk}/").data["indicadores"], 2)
        self.assertEqual(len(c.get("/api/dashboard/avance/?agrupar=criterio").data), 2)
        self.assertEqual(len(c.get("/api/dashboard/pendientes/").data), 2)
        self.assertEqual(c.get("/api/reportes/general/").data["total_indicadores"], 2)
        r = c.get("/api/reportes/evidencias/", {"estandar": self.marco.estandar_id})
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["rows"][0]["indicador"], self.marco.redaccion_oficial)
        r = c.get("/api/reportes/evidencias/exportar/?formato=xlsx")
        self.assertEqual(r.status_code, 200)
        libro = load_workbook(io.BytesIO(r.content), read_only=True)
        valores = [valor for fila in libro.active.iter_rows(values_only=True) for valor in fila]
        libro.close()
        self.assertIn(self.marco.redaccion_oficial, valores)


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ConcurrenciaAsignacionesMarcoTests(EscenarioMarco, TransactionTestCase):
    def setUp(self):
        self.crear_escenario()
        cache.clear()

    def test_creacion_duplicada_concurrente_es_controlada(self):
        self.assertEqual(connection.vendor, "postgresql")

        def crear(_):
            close_old_connections()
            try:
                user = Usuario.objects.get(pk=self.coordinador.pk)
                r = self.cliente(user).post("/api/asignaciones/", self.payload(), format="json")
                return r.status_code
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            estados = list(pool.map(crear, range(2)))
        self.assertEqual(sorted(estados), [201, 400])
        self.assertEqual(Asignacion.objects.filter(indicador_marco=self.marco).count(), 1)
        asignacion = Asignacion.objects.get(indicador_marco=self.marco)
        self.assertEqual(Auditoria.objects.filter(modelo="Asignacion", registro_id=asignacion.pk, accion="Crear asignación").count(), 1)

    def test_cambio_version_concurrente_no_deja_asignacion_incoherente(self):
        guardada = Event()
        liberar = Event()

        def cambiar_version():
            close_old_connections()
            try:
                with transaction.atomic():
                    r = self.cliente(Usuario.objects.get(pk=self.coordinador.pk)).patch(
                        f"/api/periodos/{self.periodo.pk}/", {"version_marco": self.version_b.pk}, format="json"
                    )
                    guardada.set()
                    if not liberar.wait(10):
                        raise RuntimeError("Timeout al coordinar la transacción de prueba")
                return r.status_code
            finally:
                close_old_connections()

        def crear_asignacion():
            close_old_connections()
            try:
                return self.cliente(Usuario.objects.get(pk=self.coordinador.pk)).post(
                    "/api/asignaciones/", self.payload(), format="json"
                ).status_code
            finally:
                close_old_connections()

        with ThreadPoolExecutor(max_workers=2) as pool:
            cambio = pool.submit(cambiar_version)
            try:
                self.assertTrue(guardada.wait(10))
                creacion = pool.submit(crear_asignacion)
            finally:
                liberar.set()
            self.assertEqual(cambio.result(timeout=15), 200)
            self.assertEqual(creacion.result(timeout=15), 400)
        self.periodo.refresh_from_db()
        self.assertEqual(self.periodo.version_marco_id, self.version_b.pk)
        self.assertFalse(Asignacion.objects.filter(indicador_marco__isnull=False).exists())
