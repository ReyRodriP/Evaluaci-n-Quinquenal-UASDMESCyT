"""3K: un mismo expediente recorre permisos, documentos, avisos y resultados por API."""
import csv
import io
import tempfile
import zipfile
from pathlib import Path

from django.contrib.auth.models import Group
from django.core import mail
from django.core.cache import cache
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import AccessToken

from accounts.models import Usuario
from evaluation.models import Asignacion
from evidence.models import VersionDocumento
from evidence.test_documentos import EscenarioDocumental
from notificaciones.models import Notificacion


@override_settings(AMBITOS_REVISION_EXPLICITOS=True, NOTIFICACIONES_EMAIL_ENABLED=True,
                   EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
                   PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class RecorridoIntegralTests(EscenarioDocumental, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.crear_escenario()
        cls.administrador = Usuario.objects.create_user(username='admin_3k', email='admin_3k@example.test')
        cls.administrador.groups.add(Group.objects.get_or_create(name='Administrador General')[0])

    def setUp(self):
        cache.clear()
        media = self.enterContext(tempfile.TemporaryDirectory())
        self.enterContext(override_settings(MEDIA_ROOT=Path(media)))

    def cliente(self, user=None):
        cliente = APIClient()
        cliente.credentials(HTTP_AUTHORIZATION=f'Bearer {AccessToken.for_user(user or self.coordinador)}')
        return cliente

    def comprobar(self, response, codigo=200):
        self.assertEqual(response.status_code, codigo, getattr(response, 'data', ''))
        return getattr(response, 'data', None)

    def test_recorrido_iniciado_por_coordinacion(self):
        self.recorrido(self.coordinador)

    def test_recorrido_iniciado_por_administracion(self):
        self.recorrido(self.administrador)

    def recorrido(self, creador):
        with self.captureOnCommitCallbacks(execute=True):
            a, b, aviso = self.recorrer_expediente(creador)
        destinatarios = {direccion for mensaje in mail.outbox for direccion in mensaje.to}
        self.assertIn(self.revisor.email, destinatarios)
        self.assertIn(self.responsable.email, destinatarios)
        self.assertIn(self.colaborador.email, destinatarios)
        self.assertNotIn(self.otro.email, destinatarios)
        self.assertNotIn(self.externo.email, destinatarios)
        self.comprobar_revocacion(a, b, aviso)

    def recorrer_expediente(self, creador):
        # Autenticación real JWT y alta por los dos roles globales, sin alta directa del modelo.
        self.comprobar(APIClient().get('/api/asignaciones/'), 401)
        self.comprobar(self.cliente(self.responsable).post('/api/asignaciones/', self.payload(), format='json'), 403)
        data = self.comprobar(self.cliente(creador).post('/api/asignaciones/', self.payload(), format='json'), 201)
        self.asignacion = Asignacion.objects.get(pk=data['id'])
        detalle = f'/api/asignaciones/{self.asignacion.pk}/'
        filtros = {'periodo': self.periodo.pk}
        a, b = self.solicitar('Informe integral 3K'), self.solicitar('Acta integral 3K')
        for actor in (self.responsable, self.colaborador, self.revisor):
            self.comprobar(self.cliente(actor).get(detalle, filtros))
        for actor in (self.otro, self.externo):
            self.comprobar(self.cliente(actor).get(detalle, filtros), 404)

        # Dos documentos independientes: uno aprobado y otro observado/corregido por un colaborador.
        self.aprobar_documento(a)
        self.comprobar(self.subir(b), 201)
        self.enviar(b)
        original = VersionDocumento.objects.get(solicitud_id=b)
        self.comprobar(self.revisar(b, 'solicitar_cambios', 'Completar fecha del acta 3K'))
        self.comprobar(self.cliente(self.revisor).post(detalle + 'aprobar/'), 400)
        self.comprobar(self.subir(b, self.colaborador), 201)
        self.enviar(b)
        actual = VersionDocumento.objects.get(solicitud_id=b, numero=2)
        self.comprobar(self.revisar(b, comentario='Acta corregida 3K', version=actual.pk))
        resumen = self.comprobar(self.cliente(self.revisor).get('/api/reportes/marco/', filtros))['resumen']
        self.assertEqual(resumen['documentos_aprobados'], 2)
        self.assertEqual(resumen['pendientes_cierre'], 1)
        self.assertEqual(resumen['aprobadas_globalmente'], 0)
        self.assertFalse(Notificacion.objects.filter(titulo='Asignación aprobada globalmente').exists())

        # Los avisos llevan a la versión precisa; solo los revisores reciben envíos a revisión.
        envios = Notificacion.objects.filter(titulo='Documento enviado a revisión')
        self.assertEqual(envios.count(), 3)
        self.assertEqual(set(envios.values_list('usuario_id', flat=True)), {self.revisor.pk})
        avisos = self.comprobar(self.cliente(self.revisor).get('/api/notificaciones/'))['results']
        aviso = next(n for n in avisos if n['destino']['version_id'] == actual.pk)
        self.comprobar(self.cliente(self.revisor).patch(f"/api/notificaciones/{aviso['id']}/leer/"))

        # El cierre es explícito y deja una transición global adicional, con autor y comentario.
        self.comprobar(self.cliente(self.responsable).post(detalle + 'aprobar/'), 403)
        self.comprobar(self.cliente(self.revisor).post(detalle + 'aprobar/', {'comentario': 'Cierre integral 3K'}, format='json'))
        historial = self.comprobar(self.cliente(self.responsable).get(detalle + 'historial/', filtros))
        self.assertEqual(historial['count'], 2)
        self.assertEqual(historial['results'][0]['usuario_identificador'], self.revisor.username)
        self.assertEqual(historial['results'][0]['comentario'], 'Cierre integral 3K')
        self.assertEqual(historial['results'][0]['estado_nuevo'], 'aprobado')
        self.assertEqual(Notificacion.objects.filter(titulo='Asignación aprobada globalmente').count(), 1)
        self.comprobar(self.subir(b), 400)

        # El externo solo consulta el expediente aprobado y la versión actual sin revisión interna.
        self.comprobar(self.cliente(self.externo).get(detalle, filtros))
        self.comprobar(self.cliente(self.externo).get(detalle + 'historial/', filtros), 403)
        documento = self.comprobar(self.cliente(self.externo).get(f'/api/solicitudes-documentales/{b}/'))
        self.assertEqual([v['id'] for v in documento['versiones']], [actual.pk])
        self.assertEqual(documento['versiones'][0]['revisiones'], [])
        self.comprobar(self.cliente(self.externo).get(f'/api/solicitudes-documentales/{b}/descargar/?version={original.pk}'), 404)

        # Observaciones conservadas sobre su versión original después de la corrección y cierre.
        observaciones = self.comprobar(self.cliente(self.revisor).get('/api/reportes/observaciones/', filtros))
        observacion = next(r for r in observaciones['rows'] if r['version_id'] == original.pk)
        self.assertEqual(observacion['comentario'], 'Completar fecha del acta 3K')
        resumen = self.comprobar(self.cliente(self.revisor).get('/api/reportes/marco/', filtros))['resumen']
        self.assertEqual(resumen['aprobadas_globalmente'], 1)
        self.assertEqual(resumen['pendientes_cierre'], 0)
        for reporte, formatos in (('marco', ('csv', 'xlsx', 'pdf')), ('observaciones', ('xlsx', 'pdf'))):
            for formato in formatos:
                with self.subTest(reporte=reporte, formato=formato):
                    response = self.cliente(self.revisor).get(f'/api/reportes/{reporte}/exportar/', {**filtros, 'formato': formato})
                    self.comprobar(response)
                    if formato == 'xlsx':
                        self.assertGreater(load_workbook(io.BytesIO(response.content)).active.max_row, 1)
                    elif formato == 'pdf':
                        self.assertTrue(response.content.startswith(b'%PDF'))
                    else:
                        self.assertEqual(len(list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))), 2)

        # El ZIP contiene las últimas versiones para revisión interna y consulta externa.
        for actor in (self.revisor, self.responsable, self.externo):
            response = self.cliente(actor).post('/api/evidencias/descargas-masivas/', {'solicitudes': [a, b]}, format='json')
            self.comprobar(response)
            with zipfile.ZipFile(io.BytesIO(b''.join(response.streaming_content))) as archivo:
                filas = list(csv.DictReader(io.StringIO(archivo.read('manifest.csv').decode('utf-8-sig'))))
                self.assertEqual({(r['solicitud_id'], r['version']) for r in filas}, {(str(a), '1'), (str(b), '2')})
                self.assertTrue(all(archivo.read(r['archivo']).startswith(b'%PDF') for r in filas))

        return a, b, aviso

    def comprobar_revocacion(self, a, b, aviso):
        # La revocación afecta inmediatamente al expediente, historial, avisos, reportes y ZIP.
        detalle = f'/api/asignaciones/{self.asignacion.pk}/'
        filtros = {'periodo': self.periodo.pk}
        self.ambito.activo = False
        self.ambito.save()
        for ruta in (detalle, detalle + 'historial/', '/api/reportes/marco/', '/api/reportes/observaciones/'):
            self.comprobar(self.cliente(self.revisor).get(ruta, filtros), 404)
        self.comprobar(self.cliente(self.revisor).post('/api/evidencias/descargas-masivas/', {'solicitudes': [a, b]}, format='json'), 404)
        self.comprobar(self.cliente(self.revisor).patch(f"/api/notificaciones/{aviso['id']}/leer/"), 404)
        self.assertEqual(self.comprobar(self.cliente(self.revisor).get('/api/notificaciones/'))['results'], [])
        self.comprobar(self.cliente(self.responsable).get(detalle + 'historial/', {'periodo': self.historico.pk}), 404)
